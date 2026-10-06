"""GAM media planner: an MCP server for Databricks Apps (streamable HTTP, /mcp).

Read-only against the client's Google Ad Manager network (gam.py): it finds
positions, reads typical prices from the network's own line items, checks
availability for a flight, and turns a budget into a plan (plan.py). The plan
can be saved as an Excel file to the client's results folder (a Unity Catalog
volume), written with this app's own identity.

Secrets arrive as app resources (see app.yaml): GAM_KEY_JSON, GAM_NETWORK_CODE.
"""
from __future__ import annotations

import datetime as dt
import io
import json
import os
import re

import fastmcp
import httpx

import gam
import plan as planner

mcp = fastmcp.FastMCP("GAM media planner")


@mcp.custom_route("/ping", methods=["GET"])
async def _ping(request):  # noqa: ANN001
    from starlette.responses import JSONResponse

    return JSONResponse({"status": "Healthy"})


def _day(text: str, name: str) -> dt.date:
    try:
        return dt.date.fromisoformat(str(text).strip()[:10])
    except ValueError:
        raise planner.PlanError("%s must be a date like 2026-09-01." % name)


def _fail(exc: Exception) -> dict:
    return {"ok": False, "error": str(exc)}


@mcp.tool()
def find_inventory(search: str = "", size: str = "", limit: int = 25) -> dict:
    """Find positions (ad units) in the client's Google Ad Manager network.
    search: words in the position name or path, e.g. "homepage" or "sports". size: e.g. "300x250".
    Returns each position's id, path and sizes; use the ids with check_availability and create_media_plan."""
    try:
        units = gam.ad_units()
        words = [w for w in search.lower().split() if w]
        hits = [u for u in units if all(w in u["path"].lower() for w in words) and (not size or size in u["sizes"])]
        delivered = {}
        try:
            delivered = gam.delivered_last_30_days()
        except Exception:  # noqa: BLE001 - ordering only
            pass
        hits.sort(key=lambda u: (-delivered.get(u["id"], 0), u["path"].lower()))
        limit = max(1, min(int(limit or 25), 100))
        return {"ok": True, "network": gam.network()["name"], "matches": len(hits),
                "positions": [{"ad_unit_id": u["id"], "path": u["path"], "sizes": u["sizes"],
                               "impressions_last_30_days": delivered.get(u["id"], 0)} for u in hits[:limit]]}
    except (gam.GamError, planner.PlanError) as exc:
        return _fail(exc)


@mcp.tool()
def get_rate_card() -> dict:
    """Typical prices in this network by cost method (CPM, CPC, CPD, CPV...), from its own line items.
    Use typical_rate as the default rate unless the user gives one."""
    try:
        return {"ok": True, "currency": gam.network()["currency"], "rates": gam.rate_card(),
                "note": "Median of what was actually sold in this network; the user can override any rate."}
    except gam.GamError as exc:
        return _fail(exc)


@mcp.tool()
def check_availability(ad_unit_ids: list[str], start_date: str, end_date: str, size: str = "") -> dict:
    """Available impressions for each position in a flight (dates like 2026-09-01).
    Uses Google Ad Manager's forecast; when GAM has none, estimates from the last 30 days and says so."""
    try:
        start, end = _day(start_date, "start_date"), _day(end_date, "end_date")
        by_id = {u["id"]: u for u in gam.ad_units()}
        out = []
        for uid in (ad_unit_ids or [])[:12]:  # forecasts are slow and rate-limited
            u = by_id.get(str(uid))
            if not u:
                out.append({"ad_unit_id": uid, "error": "No active position with this id."})
                continue
            sz = size or (u["sizes"][0] if u["sizes"] else "300x250")
            out.append({"ad_unit_id": uid, "path": u["path"], "size": sz, **gam.availability(uid, sz, start, end)})
        return {"ok": True, "flight": "%s to %s" % (start, end), "positions": out}
    except (gam.GamError, planner.PlanError) as exc:
        return _fail(exc)


@mcp.tool()
def create_media_plan(budget: float, start_date: str, end_date: str, ad_unit_ids: list[str] | None = None,
                      search: str = "", cost_method: str = "CPM", rate: float = 0, size: str = "",
                      strategy: str = "max_reach", max_positions: int = 6,
                      availability_override: dict | None = None, advertiser: str = "") -> dict:
    """Turn a budget into a media plan for a flight.
    Choose positions with ad_unit_ids (from find_inventory) or a search phrase; otherwise the busiest positions are used.
    cost_method: CPM, VCPM, CPV, CPC or CPD. rate: price per 1,000 for CPM/VCPM, per unit otherwise; 0 = this network's typical rate.
    strategy: max_reach (most impressions for the money) or balanced (every position gets a share).
    availability_override: {ad_unit_id: units} when the planner already knows the availability.
    Returns the plan as a table and as data for export_media_plan."""
    try:
        start, end = _day(start_date, "start_date"), _day(end_date, "end_date")
        method = cost_method.upper().strip()
        net = gam.network()
        units = gam.ad_units()
        by_id = {u["id"]: u for u in units}
        if ad_unit_ids:
            chosen = [by_id[str(i)] for i in ad_unit_ids if str(i) in by_id]
        else:
            found = find_inventory(search=search, size=size, limit=max_positions)
            if not found.get("ok"):
                return found
            chosen = [by_id[p["ad_unit_id"]] for p in found["positions"] if p["ad_unit_id"] in by_id]
        chosen = chosen[: max(1, min(int(max_positions or 6), 12))]
        if not chosen:
            raise planner.PlanError("No matching positions were found. Try another search, or list positions with find_inventory.")
        if not rate:
            card = [r for r in gam.rate_card() if r["cost_type"] == method]
            card.sort(key=lambda r: (r["line_item_type"] != "STANDARD", -r["lines"]))
            if not card:
                raise planner.PlanError("This network has no %s prices to default to. Give a rate." % method)
            rate = card[0]["typical_rate"]
        override = {str(k): int(v) for k, v in (availability_override or {}).items()}
        options = []
        for u in chosen:
            sz = size if size and size in u["sizes"] else (u["sizes"][0] if u["sizes"] else "300x250")
            if u["id"] in override:
                av = {"available": override[u["id"]], "source": "given", "note": "Availability given by the planner."}
            elif method in planner.IMPRESSION_BASED:
                av = gam.availability(u["id"], sz, start, end)
            else:
                av = {"available": 0, "source": "", "note": ""}
            options.append({"position": u["path"], "ad_unit_id": u["id"], "size": sz, "method": method, "rate": float(rate),
                            "available": av["available"], "source": av["source"]})
        result = planner.build(float(budget), start, end, options, strategy, net["currency"])
        result["advertiser"] = advertiser
        result["network"] = net["name"]
        return {"ok": True, "table": planner.to_markdown(result), "notes": result["notes"], "plan": result}
    except (gam.GamError, planner.PlanError) as exc:
        return _fail(exc)


def _m2m_token(host: str) -> str:
    r = httpx.post(host + "/oidc/v1/token", data={"grant_type": "client_credentials", "scope": "all-apis"},
                   auth=(os.environ["DATABRICKS_CLIENT_ID"], os.environ["DATABRICKS_CLIENT_SECRET"]), timeout=30)
    r.raise_for_status()
    return r.json()["access_token"]


@mcp.tool()
def export_media_plan(plan: dict, title: str = "Media plan", output_volume_dir: str = "") -> dict:
    """Save a plan from create_media_plan as an Excel file in a Unity Catalog volume folder,
    e.g. output_volume_dir="/Volumes/main/agent_portal/results". Returns the file path to share with the user."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        folder = (output_volume_dir or os.environ.get("OUTPUT_VOLUME_DIR") or "").rstrip("/")
        if not re.match(r"^/Volumes/[^/]+/[^/]+/[^/]+(/.*)?$", folder):
            raise planner.PlanError("Give a results folder like /Volumes/<catalog>/<schema>/<volume>.")
        rows = (plan or {}).get("rows") or []
        if not rows:
            raise planner.PlanError("The plan has no lines to export.")
        wb = Workbook()
        ws = wb.active
        ws.title = "Media plan"
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append(["Advertiser", plan.get("advertiser") or "", "Network", plan.get("network") or "",
                   "Budget", plan.get("budget"), "Strategy", plan.get("strategy")])
        ws.append([])
        head = ["Position", "Ad size", "Creative type", "Served by", "Flight", "Cost method", "Unit type", "Rate", "Units", "Cost", "Availability"]
        ws.append(head)
        for c in ws[4]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0B5E66")
        for r in rows:
            ws.append([r["position"], r["ad_size"], r["creative_type"], r["served_by"], r["flight"], r["cost_method"],
                       r["unit_type"], r["rate"], r["units"], r["cost"], r.get("availability_source", "")])
        ws.append(["Total", "", "", "", "", "", "", "", "", plan.get("spent")])
        ws[ws.max_row][0].font = Font(bold=True)
        for note in plan.get("notes") or []:
            ws.append([note])
        for col, width in zip("ABCDEFGHIJK", (48, 10, 14, 18, 24, 12, 18, 10, 14, 14, 14)):
            ws.column_dimensions[col].width = width
        buf = io.BytesIO()
        wb.save(buf)
        name = re.sub(r"[^A-Za-z0-9._-]+", "_", title).strip("_")[:60] or "media_plan"
        path = "%s/%s_%s.xlsx" % (folder, name, dt.datetime.now().strftime("%Y%m%d_%H%M%S"))
        host = os.environ.get("DATABRICKS_HOST", "")
        host = host if host.startswith("http") else "https://" + host
        r = httpx.put(host + "/api/2.0/fs/files" + path, params={"overwrite": "true"}, content=buf.getvalue(),
                      headers={"Authorization": "Bearer " + _m2m_token(host)}, timeout=60)
        if r.status_code >= 400:
            raise planner.PlanError("Could not save to %s (HTTP %d). The tool needs write access to that folder." % (folder, r.status_code))
        return {"ok": True, "path": path, "rows": len(rows)}
    except (planner.PlanError, KeyError) as exc:
        return _fail(exc)


if __name__ == "__main__":
    mcp.run(transport="streamable-http", host="0.0.0.0", port=int(os.environ.get("PORT", "8000")))
